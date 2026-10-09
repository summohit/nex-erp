#!/usr/bin/env python3
"""Export every task the logged-in Workway account can see, with full detail.

Workway's own Export gives one line per task. HR wanted the whole record, so
this also opens each task and reads its details panel, description and every
tab: sub-tasks, files, comments, timesheet, notes and history.

Visibility is Workway's: an employee login sees only its own tasks, an admin
login sees everyone's. Read-only — nothing is changed in Workway.

    python workway_tasks.py --storage-state auth/cestech-hr.json --out out_hr/tasks-akshara.xlsx

Writes an Excel file with two sheets: Tasks (one row per task) and Timesheet
(one row per time entry), plus the raw JSON beside it.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from workway_jobs import html_to_value  # DataTables cell flattening

log = logging.getLogger("workway.tasks")
BASE = "https://cestech.workway.pro"
TABS = [("sub_task", "Sub Tasks"), ("file", "Files"), ("comments", "Comments"),
        ("time_logs", "Timesheet"), ("notes", "Notes"), ("history", "History")]
EMPTY = re.compile(r"no record found|seems like no .* exists|no data|no (file|files) uploaded|no (note|comment|history)s? found", re.I)

# Label/value pairs from the task's details panel, and the description.
DETAILS_JS = r"""
() => {
  const clean = s => (s || '').replace(/[ \t]+/g, ' ').replace(/\n{3,}/g, '\n\n').trim();
  const labels = ['Project','Priority','Assigned To','Short Code','Milestones','Assigned By',
                  'Label','Task category','Created On','Start Date','Due Date','Estimate','Hours Logged','Status'];
  const out = {};
  const nodes = [...document.querySelectorAll('p, span, div, h6, label')];
  for (const name of labels) {
    const el = nodes.find(n => n.children.length === 0 && n.innerText && n.innerText.trim() === name);
    if (!el) continue;
    let v = el.nextElementSibling || el.parentElement.nextElementSibling;
    if (!v && el.parentElement) v = el.parentElement.parentElement && el.parentElement.parentElement.children[1];
    if (v) out[name] = clean(v.innerText).replace(/\s*It's you\s*/, ' ').trim();
  }
  const d = nodes.find(n => n.children.length === 0 && n.innerText && n.innerText.trim() === 'Description');
  let desc = '';
  if (d) {
    const v = d.nextElementSibling || d.parentElement.nextElementSibling;
    if (v) {
      const c = v.cloneNode(true);
      c.querySelectorAll('ol').forEach(ol => [...ol.children].forEach((li, i) => li.prepend((i + 1) + '. ')));
      c.querySelectorAll('ul > li').forEach(li => li.prepend('• '));
      document.body.appendChild(c); desc = clean(c.innerText); c.remove();
    }
  }
  const h = document.querySelector('h3, .heading-h3, .f-21');
  return { title: h ? clean(h.innerText) : '', fields: out, description: desc };
}
"""

# The active tab: its tables as rows, and its text.
TAB_JS = r"""
() => {
  const pane = document.querySelector('#nav-tabContent') || document.querySelector('.tab-content');
  if (!pane) return { text: '', tables: [] };
  const tables = [...pane.querySelectorAll('table')].map(tb =>
    [...tb.querySelectorAll('tr')].map(tr => [...tr.querySelectorAll('th,td')]
      .map(c => c.innerText.trim().replace(/\s+/g, ' '))));
  let text = pane.innerText.replace(/^\s*(Add (Comment|Note|File|Sub Task)|Upload File)\s*/i, '');
  return { text: text.replace(/[ \t]+/g, ' ').replace(/\n{3,}/g, '\n\n').trim(), tables };
}
"""


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--storage-state", default="auth/cestech-hr.json")
    p.add_argument("--out", default="out_hr/tasks.xlsx")
    p.add_argument("--limit", type=int, default=0, help="only the first N tasks (for a test run)")
    p.add_argument("-v", "--verbose", action="store_true")
    args = p.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(levelname)-7s %(name)s: %(message)s")

    state = Path(args.storage_state)
    if not state.exists():
        log.error("no saved session at %s — run scrape.py --login first", state)
        return 4

    from playwright.sync_api import sync_playwright

    tasks: list[dict] = []
    with sync_playwright() as pw:
        b = pw.chromium.launch(headless=True)
        ctx = b.new_context(storage_state=str(state))
        page = ctx.new_page()

        # ── the list: capture the DataTables request and ask for all of it ──
        tmpl: dict[str, str] = {}
        page.on("response", lambda r: tmpl.setdefault("url", r.url)
                if "draw=" in r.url and "/account/tasks" in r.url else None)
        page.goto(f"{BASE}/account/tasks", wait_until="networkidle", timeout=60000)
        page.wait_for_timeout(2500)
        if "/login" in page.url:
            log.error("the saved session is logged out — run scrape.py --login again")
            return 4
        url = tmpl.get("url", "")
        if not url:
            log.error("could not capture the tasks list request")
            return 1
        url = re.sub(r"([?&])length=-?\d+", r"\g<1>length=5000", url)
        # No date window and no status filter: every task, finished or not.
        url = re.sub(r"([?&])(startDate|endDate)=[^&]*", r"\g<1>\g<2>=", url)
        url = re.sub(r"([?&])status=[^&]*", r"\g<1>status=all", url)
        r = ctx.request.get(url, headers={"X-Requested-With": "XMLHttpRequest",
                                          "Accept": "application/json, text/javascript, */*; q=0.01"},
                            timeout=120000)
        if r.status != 200:
            log.error("tasks list returned HTTP %s", r.status)
            return 1
        rows = [x for x in r.json().get("data", []) if isinstance(x, dict)]
        log.info("tasks in the list: %d", len(rows))
        if args.limit:
            rows = rows[: args.limit]

        skip = {"check", "action", "DT_RowId", "DT_RowIndex", "DT_RowClass"}
        for i, row in enumerate(rows, 1):
            base = {k: html_to_value(v) for k, v in row.items()
                    if k not in skip and not isinstance(v, (dict, list))}
            tid = row.get("id")
            task = {"id": tid, "list": base, "url": f"{BASE}/account/tasks/{tid}"}
            try:
                page.goto(task["url"], wait_until="networkidle", timeout=60000)
                page.wait_for_timeout(800)
                task.update(page.evaluate(DETAILS_JS))
                task["tabs"] = {}
                for view, label in TABS:
                    page.goto(f"{task['url']}?view={view}", wait_until="networkidle", timeout=60000)
                    page.wait_for_timeout(600)
                    tab = page.evaluate(TAB_JS)
                    if EMPTY.search(tab["text"] or "") and len(tab["text"]) < 160:
                        tab = {"text": "", "tables": []}
                    task["tabs"][label] = tab
            except Exception as e:  # one bad page must not lose the rest
                log.warning("task %s: %s", tid, e)
                task["error"] = str(e)
            task["title"] = re.sub(r"\s+(Low|Medium|High|Urgent)$", "", base.get("heading", "")).strip() or task.get("title", "")
            tasks.append(task)
            log.info("[%d/%d] %s", i, len(rows), task.get("title") or base.get("heading", tid))
        b.close()

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.with_suffix(".json").write_text(json.dumps(tasks, ensure_ascii=False, indent=2), encoding="utf-8")
    write_xlsx(out, tasks)
    log.info("done: %d task(s) -> %s", len(tasks), out)
    return 0


def write_xlsx(path: Path, tasks: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    field_names = ["Project", "Priority", "Assigned To", "Assigned By", "Short Code", "Milestones",
                   "Label", "Task category"]
    head = (["Task ID", "Task", "Status", "Created On", "Start Date", "Due Date", "Completed On", "Hours Logged"]
            + field_names + ["Description"] + [label for _, label in TABS] + ["Link"])
    ws.append(head)
    for t in tasks:
        L, F = t.get("list", {}), t.get("fields", {})
        status = L.get("board_column") or L.get("status") or ""
        ws.append([
            t.get("id"), t.get("title") or L.get("heading", ""), status,
            F.get("Created On", ""), L.get("start_date", ""), L.get("due_date", ""), L.get("completed_on", ""),
            F.get("Hours Logged") or L.get("timeLogged") or L.get("time_logged", ""),
            *[F.get(n, "") for n in field_names],
            t.get("description", ""),
            *[(t.get("tabs", {}).get(label, {}) or {}).get("text", "") for _, label in TABS],
            t.get("url", ""),
        ])

    ts = wb.create_sheet("Timesheet")
    ts.append(["Task ID", "Task", "Employee", "Start Time", "End Time", "Memo", "Hours Logged"])
    for t in tasks:
        for table in (t.get("tabs", {}).get("Timesheet", {}) or {}).get("tables", []):
            for row in table[1:]:
                if len(row) >= 5:
                    ts.append([t.get("id"), t.get("title", ""), *row[:5]])

    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1373E5")
    for sheet in (ws, ts):
        for c in sheet[1]:
            c.font, c.fill = bold, fill
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = min(60, max(10, *(len(str(c.value or "").split("\n")[0]) for c in col)))
            sheet.column_dimensions[col[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


if __name__ == "__main__":
    sys.exit(main())
